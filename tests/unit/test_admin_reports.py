"""Report files: untrusted email text must never become a spreadsheet formula."""

import csv
import io
from datetime import UTC, datetime

from openpyxl import load_workbook  # type: ignore[import-untyped]

from ci_admin.reports import label_title, parse, table, to_csv, to_xlsx

HOSTILE = ['=HYPERLINK("https://evil.example")', "+1+1", "-2+3", "@SUM(A1)", "\tx"]


def report(rows: list[dict[str, object]]) -> dict[str, object]:
    return {
        "dataset": "messages",
        "group_by": None,
        "label_mode": "names",
        "columns": [
            {"key": "received_at", "title": "Received", "kind": "datetime"},
            {"key": "subject", "title": "Subject", "kind": "text"},
            {"key": "labels", "title": "Labels", "kind": "labels"},
        ],
        "rows": rows,
    }


def rows() -> list[dict[str, object]]:
    when = datetime(2026, 10, 28, 15, 30, tzinfo=UTC).isoformat()
    return [{"received_at": when, "subject": s, "labels": ["cat_payment_billing"]} for s in HOSTILE]


def test_csv_neutralizes_formulas_and_names_labels() -> None:
    header, body = table(report(rows()), "America/Chicago")
    text = to_csv(header, body).decode("utf-8-sig")
    parsed = list(csv.reader(io.StringIO(text)))
    assert parsed[0] == ["Received", "Subject", "Labels"]
    assert parsed[1][0] == "2026-10-28 10:30"  # shown in the console's time zone
    assert [r[1] for r in parsed[1:]] == ["'" + s for s in HOSTILE]
    assert {r[2] for r in parsed[1:]} == {"Payment"}


def test_xlsx_cells_are_strings_not_formulas() -> None:
    header, body = table(report(rows()), "America/Chicago")
    wb = load_workbook(io.BytesIO(to_xlsx(header, body, "Messages")))
    ws = wb.active
    cells = [ws.cell(row=i, column=2) for i in range(2, 2 + len(HOSTILE))]
    assert [c.value for c in cells] == HOSTILE
    assert all(c.data_type == "s" for c in cells)
    assert ws.cell(row=2, column=1).is_date


def test_keys_mode_and_column_choice() -> None:
    req = parse({"dataset": "messages", "columns": ["subject", "nope"], "label_mode": "keys"})
    assert req.columns == ("subject",) and req.label_mode == "keys"
    assert parse({"dataset": "messages"}).columns[0] == "received_at"
    assert label_title("cat_payment_billing") == "Payment"
    assert label_title("cat_new_thing") == "New thing"
